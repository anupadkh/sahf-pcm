"""Read and validate v2 workbooks without assuming a fixed active worksheet."""
import io
import zipfile
import re
from datetime import datetime
from openpyxl import load_workbook
from fastapi import HTTPException
from .schemas import Project
from .utilities_functions import amount, cell_date

def open_workbook(content):
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 50*1024*1024:
                raise ValueError('Workbook is too large')
        return load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(422,'Unreadable or oversized XLSX workbook')

def normalized(value):
    return ' '.join(str(value or '').split()).casefold()

def header_map(row):
    names={normalized(v):i for i,v in enumerate(row) if v is not None}
    return names if 'project title' in names and 'commitment amount (usd)' in names else None

def discover(workbook):
    sheets=[]
    for sheet in workbook:
        header=None; count=0; header_row=0
        for number,row in enumerate(sheet.iter_rows(values_only=True),1):
            if header is None and number<=10:
                header=header_map(row)
                if header: header_row=number
            elif header and row[header['project title']] and not normalized(row[header['project title']]).startswith('input the title'):
                count+=1
        if header:sheets.append({'name':sheet.title,'rows':count,'header_row':header_row})
    return sheets

def choose_sheet(sheets):
    populated=[s for s in sheets if s['rows']]
    for name in ['PCM Template','(Rev) ADB and WBG Proj']:
        if any(s['name']==name for s in populated):return name
    return populated[0]['name'] if populated else sheets[0]['name'] if sheets else None

def text(value):return str(value or '').strip()
def flag(value, field):
    value=normalized(value)
    if value not in ['yes','no']:raise ValueError(f'{field} must be Yes or No')
    return value=='yes'
def enum(value):return text(value).capitalize()

def funding_amount(value):
    raw=text(value)
    if normalized(value) in ['not disclosed (ta)','not disclosed','n/a','unknown','']:
        return None
    match=re.match(r'^\s*(?:US\$|USD|\$)?\s*~?\s*(\d[\d,]*(?:\.\d+)?)\s*(million|billion)?(?=\s|$)',raw,re.I)
    if match and (match[2] or raw.startswith(('US$','USD','$'))):
        return float(match[1].replace(',',''))*{'million':1e6,'billion':1e9,None:1}.get(match[2].lower() if match[2] else None,1)
    return amount(value)

def parse_rows(sheet,user,header_row):
    from .config import TAXONOMY
    from pydantic import ValidationError
    rows=sheet.iter_rows(values_only=True)
    for _ in range(header_row-1):next(rows)
    headers=next(rows);mapping=header_map(headers)
    parsed=[];errors=[];warnings=[]
    for index,row in enumerate(rows,header_row+1):
        if not any(v is not None for v in row):continue
        raw_title=row[mapping['project title']] if mapping['project title']<len(row) else None
        if not raw_title or normalized(raw_title).startswith('input the title'):continue
        if all(normalized(v)=='guidance' for v in row if v is not None):continue
        issues=[];sources={};labels={};data={};fallbacks={}
        def issue(label,value):
            entry=f'{label}: {value if value is not None else "(empty)"}'
            if entry not in issues:issues.append(entry)
        def read(label,aliases=()):
            col=next((mapping[normalized(n)] for n in (label,*aliases) if normalized(n) in mapping),None)
            return row[col] if col is not None and col<len(row) else None
        specs=[('title','Project Title',text,'Untitled imported project',()),('description','Project Details (Short Description)',text,'',('Description',)),('external_id','Project ID',text,'',()),('implementing_partner','Implementing Partner',text,'',()),('funding_partner','Funding Partner',text,'Unknown',()),('country','Country',text,user.get('country') if user.get('country') in TAXONOMY['countries'] else TAXONOMY['countries'][0],()),('project_status','Project Status',enum,'Ongoing',()),('start_date','Start Date',cell_date,None,()),('end_date','End Date',cell_date,None,()),('commitment_usd','Commitment Amount (USD)',funding_amount,None,()),('ibf_relevance','IBF Relevance',enum,'No',()),('hazard','Hazard Type',text,'Flood',()),('sector','Sector',text,'Unspecified',()),('multi_sector','Multi-sector?',lambda v:flag(v,'Multi-sector'),False,()),('notes','Notes / Remarks',text,'',('Notes',))]
        for field,label,convert,fallback,aliases in specs:
            raw=read(label,aliases);sources[field]=raw;labels[field]=label;fallbacks[field]=fallback
            try:data[field]=convert(raw)
            except Exception:issue(label,raw);data[field]=fallback
        for field,options in [('country',TAXONOMY['countries']),('hazard',TAXONOMY['hazards']),('sector',TAXONOMY['sectors'])]:
            if data[field] not in options:
                issue(labels[field],sources[field]);data[field]=fallbacks[field]
                if field=='sector' and len(text(sources[field]))>80:data['notes']+='\nOriginal Sector cell (requires review): '+text(sources[field])
        data['pillars']=[]
        for pillar in range(1,5):
            col=next((i for i,h in enumerate(headers) if normalized(h).startswith(f'ew pillar {pillar}:') or normalized(h)==f'ew pillar {pillar}'),None)
            label=text(headers[col]) if col is not None else f'EW Pillar {pillar}'
            raw=row[col] if col is not None and col<len(row) else None
            try:
                if flag(raw,label):data['pillars'].append(pillar)
            except ValueError:issue(label,raw)
        raw_amount=sources['commitment_usd']
        if isinstance(raw_amount,str):data['notes']+='\nOriginal commitment: '+raw_amount
        data['confidence']='Low' if issues or data['commitment_usd'] is None or '~' in text(raw_amount) else 'Medium'
        # Collect every field error, preserving the source value before using an editable draft fallback.
        for _ in range(3):
            try:
                project=Project(**data);break
            except ValidationError as exc:
                for error in exc.errors():
                    field=error['loc'][0] if error['loc'] else None
                    if field in fallbacks:
                        issue(labels[field],sources[field]);data[field]=fallbacks[field]
                    else:
                        issue('Start Date',sources['start_date']);issue('End Date',sources['end_date']);data['start_date']=None;data['end_date']=None
        else:
            errors.append({'row':index,'message':'Row could not be converted into an editable draft'});continue
        project.import_issues='\n'.join(issues)
        if issues:project.confidence='Low'
        if issues:warnings.append({'row':index,'message':'Issues while importing from excel: '+ '; '.join(issues)})
        elif data['commitment_usd'] is None or '~' in text(raw_amount) or 'indicated' in text(raw_amount).lower():warnings.append({'row':index,'message':'Commitment is undisclosed or approximate; source text retained for review.'})
        parsed.append((index,project))
    return parsed,errors,0,warnings

